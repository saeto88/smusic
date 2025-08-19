from docx import Document

doc = Document('Protipo webapp musica .docx')

with open('prototipo_webapp_musica.txt', 'w', encoding='utf-8') as f:
    for paragraph in doc.paragraphs:
        f.write(paragraph.text + '\n')
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                f.write(cell.text + '\t')
            f.write('\n')

print("Document converted to prototipo_webapp_musica.txt")